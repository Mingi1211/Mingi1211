import sys, os, re
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, 'mathlib'))
import latex2mathml.converter as l2m
from lxml import etree
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1]
FONT = '바탕체'
XSL = etree.XSLT(etree.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))
doc = Document()

st = doc.styles['Normal']
st.font.name = FONT
st.font.size = Pt(10)
rpr = st.element.get_or_add_rPr()
f = rpr.find(qn('w:rFonts'))
if f is None:
    f = OxmlElement('w:rFonts'); rpr.insert(0, f)
for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
    f.set(qn(k), FONT)
st.paragraph_format.space_after = Pt(2)
st.paragraph_format.line_spacing = 1.2
st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.0)


def omml(tex):
    tex = re.sub(r'(\\begin\{bmatrix\}.*?\\end\{bmatrix\})', r'{\1}', tex)   # 행렬마다 괄호가 따로 늘어나게
    tex = tex.replace(r'\qquad', r'\text{      }').replace(r'\quad', r'\text{   }')   # mspace는 워드에서 사라짐
    mml = l2m.convert(tex).replace('<mover>', '<mover accent="true">')   # 점·물결 기호를 악센트로
    node = XSL(etree.fromstring(mml.encode())).getroot()
    # 수식 글자 크기 10pt
    W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    for r in node.iter('{%s}r' % M):
        wr = etree.SubElement(r, '{%s}rPr' % W)
        sz = etree.SubElement(wr, '{%s}sz' % W); sz.set('{%s}val' % W, '20')
        r.remove(wr); r.insert(1 if r.find('{%s}rPr' % M) is not None else 0, wr)
    return node


def _run(p, text, bold=False, size=10):
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = FONT
    rf = r._element.get_or_add_rPr()
    ff = rf.find(qn('w:rFonts'))
    if ff is None:
        ff = OxmlElement('w:rFonts'); rf.insert(0, ff)
    for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        ff.set(qn(k), FONT)
    return r


def _rich(p, text, size=10):
    """$...$ = 인라인 수식, **...** = 굵게 (수식을 사이에 둔 굵게 범위도 유지)"""
    bold = False
    for tok in re.split(r'(\$[^$]*\$|\*\*)', text):
        if not tok:
            continue
        if tok == '**':
            bold = not bold
        elif tok.startswith('$') and tok.endswith('$') and len(tok) > 1:
            p._p.append(omml(tok[1:-1]))
        else:
            _run(p, tok, bold=bold, size=size)


def title(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(10)
    _run(p, text, bold=True, size=13)


def chap(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    _run(p, text, bold=True, size=11)


def sub(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    _run(p, text, bold=True, size=10)


def b(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.45); p.paragraph_format.first_line_indent = Cm(-0.35)
    _rich(p, '- ' + text)


def eq(tex):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(4)
    M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    para = etree.SubElement(p._p, '{%s}oMathPara' % M)
    para.append(omml(tex))


def cy(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.45); p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_before = Pt(2)
    _rich(p, '→ **실습에서**: ' + text)


# =====================================================================
title('Cyclo 실습 MR 개념 정리')

# ---------------------------------------------------------------- 3.2.3
chap('3.2.3 회전의 지수 좌표')
b('어떤 자세든 기준 자세에서 **단위축 $\\hat{\\omega}$ 둘레로 $\\theta$만큼 돌리면** 만들 수 있다. $\\hat{\\omega}\\theta$(3-벡터)를 **지수 좌표**라 하고, 회전행렬 대신 쓰는 3개짜리 자세 표현이다.')
b('"지수"인 이유: $\\hat{\\omega}$를 각속도로 보고 $\\theta$초 동안 적분하는 문제이기 때문. 스칼라 $\\dot{x}=ax$의 해가 $x(t)=e^{at}x(0)$이듯, 벡터 $\\dot{x}=Ax$의 해는 $x(t)=e^{At}x(0)$ (**행렬 지수**).')
b('회전하는 벡터 $p$는 $\\dot{p}=\\hat{\\omega}\\times p=[\\hat{\\omega}]p$를 만족하므로 $p(\\theta)=e^{[\\hat{\\omega}]\\theta}p(0)$. 3×3 반대칭 행렬에서는 급수가 닫힌 꼴이 된다(**Rodrigues 공식**).')
eq(r'\mathrm{Rot}(\hat{\omega},\theta)=e^{[\hat{\omega}]\theta}=I+\sin\theta\,[\hat{\omega}]+(1-\cos\theta)\,[\hat{\omega}]^{2}')
b('**exp**: 각속도 $\\hat{\\omega}$를 $\\theta$초 적분해 $I$에서 $R$로 간다(적분과 비슷). **log**: $R$을 만드는 단위 각속도와 시간을 돌려준다(미분과 비슷). log는 Rodrigues 공식을 거꾸로 푸는 알고리즘이다.')
b('회전 관절에서는 $\\hat{\\omega}$ = 관절축, $\\theta$ = 관절각.')
cy('자세 보간 $R(s)=R_{start}\\exp(\\log(R_{start}^{T}R_{end})\\,s)$ (rotationCubic), 자세 오차 $\\log(R_{d}R^{T})$ (computeDesiredVelocity).')

# ---------------------------------------------------------------- 5.1
chap('5.1 자코비안')
b('자코비안 $J$는 관절 속도를 말단 속도로 바꾼다: $\\mathcal{V}=J(\\theta)\\dot{\\theta}$. **$i$번째 열 = 관절 $i$만 속도 1, 나머지 0일 때의 말단 속도.** 관절 $n$개면 $6\\times n$.')
sub('5.1.1 공간 자코비안 J_s')
b('말단 속도를 공간 좌표계 {s}에서 본 twist로 쓴다: $\\mathcal{V}_{s}=J_{s}(\\theta)\\dot{\\theta}$.')
b('관절 $i$의 열은 **관절 $i$와 {s} 사이에 있는 관절($1\\sim i-1$)의 각도에만** 영향을 받는다. 그 뒤 관절이 움직여도 관절 $i$와 {s}의 관계는 그대로이기 때문.')
eq(r'J_{s1}=\mathcal{S}_{1},\qquad J_{si}=\left[\mathrm{Ad}_{e^{[\mathcal{S}_{1}]\theta_{1}}\cdots e^{[\mathcal{S}_{i-1}]\theta_{i-1}}}\right]\mathcal{S}_{i}')
b('$\\mathcal{S}_{i}$는 영점 자세에서 {s}로 본 관절 $i$의 스크루 축. 미분이 필요 없고, 말단 좌표계 {b}를 어떻게 잡든 상관없다.')
sub('5.1.2 물체 자코비안 J_b')
b('말단 좌표계 {b}에서 본 twist로 쓴다: $\\mathcal{V}_{b}=J_{b}(\\theta)\\dot{\\theta}$. $\\mathcal{B}_{i}$는 영점 자세에서 {b}로 본 관절 $i$의 스크루 축.')
b('관절 $i$의 열은 **관절 $i$와 {b} 사이 관절($i+1\\sim n$)에만** 의존한다. 마지막 열 $J_{bn}=\\mathcal{B}_{n}$. {s}를 어떻게 잡든 상관없다.')
eq(r'J_{bi}=\left[\mathrm{Ad}_{e^{-[\mathcal{B}_{n}]\theta_{n}}\cdots e^{-[\mathcal{B}_{i+1}]\theta_{i+1}}}\right]\mathcal{B}_{i}')
b('각 열이 twist이므로 좌표계를 바꾸는 규칙으로 서로 변환된다.')
eq(r'J_{b}=\left[\mathrm{Ad}_{T_{bs}}\right]J_{s},\qquad J_{s}=\left[\mathrm{Ad}_{T_{sb}}\right]J_{b}')
cy('Cyclo(Pinocchio)는 LOCAL_WORLD_ALIGNED — 말단 원점의 속도를 world 축으로 쓴다. 기준점은 {b}, 축은 {s}라서 $J_{s}$도 $J_{b}$도 아니다.')

# ---------------------------------------------------------------- 5.3
chap('5.3 특이점')
b('자코비안의 두 용도: 관절 속도 → 말단 twist ($\\mathcal{V}=J\\dot{\\theta}$), 말단 렌치 → 관절 토크 ($\\tau=J^{T}\\mathcal{F}$).')
b('$J$는 $6\\times n$이라 $\\mathrm{rank}\\,J\\le\\min(6,n)$. rank가 $\\min(6,n)$이면 **full rank**, 어떤 자세 $\\theta^{*}$에서 rank가 최대값보다 작아지면 그 자세는 **특이(singular)**. 특이 자세에서는 **한 방향 이상으로 움직일 능력을 잃는다.**')
b('관절 수로 분류: $n<6$ → $J$가 세로로 김, **운동학적으로 부족**(말단이 임의의 6차원 운동 불가) / $n=6$ → 정방, 범용 매니퓰레이터 / $n>6$ → $J$가 가로로 김, **여유(redundant)**. 같은 말단 twist를 다른 관절 속도로 낼 수 있어 말단에 보이지 않는 **내부 운동**이 가능 (손을 고정한 채 팔꿈치를 움직이는 것).')
b('역문제 "$v_{tip}$이 주어지면 $\\dot{\\theta}$는?"는 $J$가 정방이 아니거나 특이라 역행렬이 없어 쉽지 않다. 여유 로봇은 해가 무수히 많다 → 6장.')
b('평면 3R 팔을 쭉 펴면 rank 2 → 1: 모든 관절이 끝을 수직으로만 움직이고 **수평 속도는 불가**. 대신 수평 힘은 관절 토크 없이 **구조가 수동적으로** 버틴다.')
cy('AI Worker 팔은 7관절이라 $6\\times 7$(여유 1자유도). C5에서 팔을 끝까지 뻗은 자세가 작업공간 경계 특이점이고, Cyclo는 감쇠 항으로 관절 속도 폭주를 막았다.')

# ---------------------------------------------------------------- 6.2
chap('6.2 수치 역기구학 (+ 6.3 역속도 기구학)')
b('역기구학 = $x_{d}-f(\\theta_{d})=0$을 만족하는 $\\theta_{d}$ 찾기 → 근 찾기 문제 → **Newton–Raphson**.')
b('1차원 그림: 현재 추정 $\\theta_{0}$에서 기울기(접선)를 그어 $\\theta$축과 만나는 점을 다음 추정 $\\theta_{1}$로 삼고 반복. 함수가 선형이면 한 번에 정확. 해가 여러 개면 **초기값 근처의 해**로 수렴하고, 기울기가 작은 평평한 곳 근처에서 시작하면 멀리 튀어 수렴이 어렵다 → 초기값이 해와 가까워야 한다.')
b('벡터로 일반화: 테일러 1차 근사로 $x_{d}-f(\\theta_{i})=J(\\theta_{i})\\Delta\\theta$ → $\\Delta\\theta=J^{-1}(x_{d}-f(\\theta_{i}))$. $J$가 정방이 아니거나 특이하면 **의사역행렬 $J^{\\dagger}$**를 쓴다.')
b('$J^{\\dagger}$의 성질: 해가 여러 개면(여유 로봇) **길이가 가장 짧은 $\\Delta\\theta$**, 정확한 해가 없으면(특이점, 관절 부족) **오차가 가장 작은 $\\Delta\\theta$**를 준다.')
eq(r'\theta_{i+1}=\theta_{i}+J^{\dagger}(\theta_{i})\,e,\qquad e=x_{d}-f(\theta_{i})')
b('$e$가 충분히 작으면 종료. 제어기 안에서 쓸 때는 매 시간 갱신된 $x_{d}$에 대해 풀고, **직전 해를 초기 추정**으로 쓴다(목표가 조금씩만 바뀌므로).')
b('목표가 변환행렬 $T_{sd}$일 때: 오차 $e$를 **"단위 시간 동안 따라가면 목표에 도착하는 twist"**로 바꾼다.')
eq(r'T_{bd}=T_{sb}^{-1}T_{sd},\qquad [\mathcal{V}_{b}]=\log(T_{bd}),\qquad \theta_{i+1}=\theta_{i}+J_{b}^{\dagger}(\theta_{i})\,\mathcal{V}_{b}')
b('$\\omega_{b}$와 $v_{b}$가 충분히 작으면 종료.')
b('**역속도 기구학**: 원하는 twist $\\mathcal{V}_{d}$를 내는 관절 속도 $\\dot{\\theta}=J^{\\dagger}\\mathcal{V}_{d}$ ($J$와 $\\mathcal{V}_{d}$는 같은 좌표계). $n>6$이면 $\\dot{\\theta}$ 성분 제곱합이 최소, $n<6$이면 twist 오차 제곱합이 최소인 해. 다른 양을 최소화하는 역행렬(가중)도 있다.')
cy('Cyclo는 Newton–Raphson을 수렴할 때까지 돌리지 않고, **매 10 ms 역속도 기구학을 QP(가중치 + 감쇠 + 제약)로 한 번 풀고 적분**한다. $k_{p}\\Delta t=50\\times 0.01=0.5$라서 "한 주기에 오차를 절반 줄이는 NR 한 걸음"과 비슷하다.')

# ---------------------------------------------------------------- 8.3
chap('8.3 뉴턴–오일러 역동역학 (RNEA)')
b('**입력** 관절 위치·속도·가속도 $\\theta,\\dot{\\theta},\\ddot{\\theta}$와 말단이 환경에 가하는 렌치 $\\mathcal{F}_{tip}$ → **출력** 필요한 관절 토크 $\\tau$.')
b('링크 $i$의 질량중심에 좌표계 {i}, 말단 {n+1}, 월드 {0}. $\\mathcal{V}_{i}$ = 링크 $i$의 twist({i} 기준), $\\mathcal{A}_{i}$ = {i}에서 본 관절 $i$의 스크루 축, $M_{i,i-1}$ = 관절 $i$가 0일 때 {i}에서 본 {i−1}.')
b('**중력**: 베이스 가속도 $\\dot{\\mathcal{V}}_{0}$를 중력 반대 방향의 선가속도로 둔다. 중력은 위로 가속하는 것과 구별할 수 없기 때문.')
sub('앞으로 (링크 1 → n): 각 링크의 위치·속도·가속도')
eq(r'T_{i,i-1}=e^{-[\mathcal{A}_{i}]\theta_{i}}M_{i,i-1}')
eq(r'\mathcal{V}_{i}=\left[\mathrm{Ad}_{T_{i,i-1}}\right]\mathcal{V}_{i-1}+\mathcal{A}_{i}\dot{\theta}_{i}')
eq(r'\dot{\mathcal{V}}_{i}=\left[\mathrm{Ad}_{T_{i,i-1}}\right]\dot{\mathcal{V}}_{i-1}+\left[\mathrm{ad}_{\mathcal{V}_{i}}\right]\mathcal{A}_{i}\dot{\theta}_{i}+\mathcal{A}_{i}\ddot{\theta}_{i}')
b('앞 링크의 속도를 내 좌표계로 옮기고, 내 관절이 더하는 속도를 합친다. 가속도에는 속도끼리 곱한 항이 붙는다.')
sub('뒤로 (링크 n → 1): 각 관절이 내야 할 힘·토크')
eq(r'\mathcal{F}_{i}=\left[\mathrm{Ad}_{T_{i+1,i}}\right]^{T}\mathcal{F}_{i+1}+\mathcal{G}_{i}\dot{\mathcal{V}}_{i}-\left[\mathrm{ad}_{\mathcal{V}_{i}}\right]^{T}\left(\mathcal{G}_{i}\mathcal{V}_{i}\right),\qquad \mathcal{F}_{n+1}=\mathcal{F}_{tip}')
eq(r'\tau_{i}=\mathcal{F}_{i}^{T}\mathcal{A}_{i}')
b('다음 링크가 필요한 렌치를 내 좌표계로 옮기고, 내 링크를 가속하는 데 필요한 렌치(강체 역동역학)를 더한다. 그중 **관절 축 방향 성분만 모터가** 내고, 나머지는 베어링 같은 구조가 버틴다.')
b('장점: 미분이 필요 없고, 재귀라서 계산이 효율적이다.')
cy('C2의 qfrc_bias = MuJoCo가 이 계산을 $\\ddot{\\theta}=0$으로 돌린 값 = $c(\\theta,\\dot{\\theta})+g(\\theta)$. 이걸 팔 관절에 더해 중력을 상쇄했다.')

# ---------------------------------------------------------------- 9.1–9.2
chap('9.1 – 9.2 점대점 궤적')
b('**궤적** $\\theta(t),\\ t\\in[0,T]$ = **경로** $\\theta(s),\\ s\\in[0,1]$ + **시간 스케일링** $s(t)$. 경로는 어디로 갈지, 시간 스케일링은 얼마나 빨리 갈지.')
eq(r'\dot{\theta}=\frac{d\theta}{ds}\dot{s},\qquad \ddot{\theta}=\frac{d\theta}{ds}\ddot{s}+\frac{d^{2}\theta}{ds^{2}}\dot{s}^{2}')
b('동역학이 $\\ddot{\\theta}$에 의존하므로 $\\theta(s)$와 $s(t)$ 모두 **2번 미분 가능**해야 한다.')
sub('경로')
b('**관절 공간 직선** $\\theta(s)=\\theta_{start}+s(\\theta_{end}-\\theta_{start})$: 말단은 직선으로 움직이지 않는다. 관절 한계가 서로 독립이면 가능한 관절 영역이 볼록이라 직선이 항상 그 안에 있다.')
b('**작업 공간 직선**: 점마다 역기구학이 필요하고, 직선이 작업 공간 밖을 지나면 실행할 수 없다.')
b('**SE(3)에서 두 자세 사이**: 원소끼리 빼는 건 의미가 없다. 일정한 twist를 따라가는 **스크루 경로**. twist가 {start} 기준이라 exp를 오른쪽에 곱한다.')
eq(r'X(s)=X_{start}\exp\!\left(\log\!\left(X_{start}^{-1}X_{end}\right)s\right)')
b('**위치·자세 분리 경로**: 원점은 직선, 자세는 자세끼리의 스크루 보간.')
eq(r'p(s)=p_{start}+s\,(p_{end}-p_{start}),\qquad R(s)=R_{start}\exp\!\left(\log\!\left(R_{start}^{T}R_{end}\right)s\right)')
sub('시간 스케일링')
b('**3차 다항식**: 조건 4개 $s(0)=0,\\ s(T)=1,\\ \\dot{s}(0)=\\dot{s}(T)=0$.')
eq(r's(t)=3\left(\frac{t}{T}\right)^{2}-2\left(\frac{t}{T}\right)^{3}')
b('$\\dot{s}$는 시작·끝에서 0이지만 **$\\ddot{s}$는 $t=0$에서 $6/T^{2}$로 불연속** 점프한다.')
b('**5차 다항식**: 계수 2개가 늘어 $\\ddot{s}(0)=\\ddot{s}(T)=0$ 조건까지 맞춘다 → 시작·끝 가속도 0, 더 부드럽다.')
b('**사다리꼴**: 등가속 → 등속 → 등감속. 3차처럼 가속도가 불연속. **S-curve**: 7구간(일정 jerk 포함), 시작·끝 가속도 0.')
cy('Cyclo MoveL은 **위치·자세 분리 경로 + 3차 스케일링**이다. 5차 함수는 Cyclo에 없다.')

# ---------------------------------------------------------------- 11.3
chap('11.3 속도 입력 운동 제어')
b('관절 **속도를 직접 명령**할 수 있다고 가정한다. 하위 제어기가 요청한 속도를 잘 낸다고 믿을 수 있을 때, 바퀴 로봇처럼 상위가 속도만 명령할 때 쓴다.')
b('**개루프(앞먹임)**: $\\dot{\\theta}=\\dot{\\theta}_{d}$. 위치를 재지 않으므로 오차가 생기면 회복하지 못한다.')
b('**P 제어**: $\\dot{\\theta}=K_{p}\\theta_{e}\\ (K_{p}>0)$. 설정점 제어면 $\\dot{\\theta}_{e}=-K_{p}\\theta_{e}$ → 시정수 $1/K_{p}$, $K_{p}$가 클수록 빠르다. 너무 크면 진동하고, 액추에이터 속도 한계에 걸려 선형 모델이 깨진다.')
b('일정 속도 $c$로 움직이는 목표를 P로 쫓으면 **정상상태 오차 $c/K_{p}$**가 남는다. P 제어는 움직이려면 오차가 있어야 하기 때문.')
b('**PI 제어**: $\\dot{\\theta}=K_{p}\\theta_{e}+K_{i}\\int\\theta_{e}\\,dt$. 오차 동역학이 2차 → 등속 궤적에서도 정상상태 오차 0.')
eq(r'\ddot{\theta}_{e}+K_{p}\dot{\theta}_{e}+K_{i}\theta_{e}=0,\qquad \zeta=\frac{K_{p}}{2\sqrt{K_{i}}},\qquad \omega_{n}=\sqrt{K_{i}}')
b('**임계 감쇠**($K_{i}=K_{p}^{2}/4$)가 빠르고 오버슈트가 없어 좋다.')
b('오차가 쌓일 때까지 기다리지 않도록 앞먹임을 더한 것이 **최종 형태**. 다관절이면 벡터, 게인은 양수 × 단위행렬.')
eq(r'\dot{\theta}=\dot{\theta}_{d}+K_{p}\theta_{e}+K_{i}\int\theta_{e}\,dt')
sub('11.3.3 작업 공간 운동 제어')
b('목표가 말단 자세 $X_{d}(t)$로 주어질 때. 목표 twist $[\\mathcal{V}_{d}]=X_{d}^{-1}\\dot{X}_{d}$({d} 기준), 실제 twist $\\mathcal{V}_{b}$({b} 기준).')
eq(r'\mathcal{V}_{b}=\left[\mathrm{Ad}_{X^{-1}X_{d}}\right]\mathcal{V}_{d}+K_{p}X_{e}+K_{i}\int X_{e}\,dt,\qquad [X_{e}]=\log\!\left(X^{-1}X_{d}\right)')
b('앞먹임: $\\mathcal{V}_{d}$를 **지금 말단 좌표계로 옮겨** 쓴다($X^{-1}X_{d}$ = 지금 말단에서 본 목표 자세). 오차 $X_{e}$: 빼기 대신 **지금 → 목표로 단위 시간에 가는 twist**. 관절 속도는 $\\dot{\\theta}=J_{b}^{\\dagger}\\mathcal{V}_{b}$.')
b('**분리형**: 자세 $R$과 위치 $p$를 따로. 각속도 앞먹임은 지금 말단 기준으로 돌리고, 선속도 앞먹임은 그대로.')
eq(r'\begin{bmatrix}\omega_{b}\\ \dot{p}\end{bmatrix}=\begin{bmatrix}R^{T}R_{d}&0\\ 0&I\end{bmatrix}\begin{bmatrix}\omega_{d}\\ \dot{p}_{d}\end{bmatrix}+K_{p}X_{e}+K_{i}\int X_{e}\,dt,\qquad X_{e}=\begin{bmatrix}\log(R^{T}R_{d})\\ p_{d}-p\end{bmatrix}')
b('정지한 목표로 갈 때(앞먹임 0, $K_{i}=0$) 결합형은 **일정한 스크루 축을 따라** 움직이고, 분리형은 **원점이 직선**으로 간다.')
cy('Cyclo computeDesiredVelocity는 분리형($K_{i}=0,\\ K_{p}=50$)이고, 자세 오차를 world 기준 $\\log(R_{d}R^{T})$로 쓴다.')

# ---------------------------------------------------------------- 11.4
chap('11.4 토크·힘 입력 운동 제어')
b('관절 **토크를 명령**하므로 동역학을 고려해야 한다. 1관절 로봇:')
eq(r'\tau=M\ddot{\theta}+mgr\cos\theta+b\dot{\theta}=M\ddot{\theta}+h(\theta,\dot{\theta})')
b('**PID**: 속도는 보통 엔코더 위치를 수치 미분해 얻는다. $K_{i}=0$이면 **PD**.')
eq(r'\tau=K_{p}\theta_{e}+K_{i}\int\theta_{e}\,dt+K_{d}\dot{\theta}_{e}')
sub('중력이 없을 때 PD')
eq(r'M\ddot{\theta}_{e}+(b+K_{d})\dot{\theta}_{e}+K_{p}\theta_{e}=0,\qquad \zeta=\frac{b+K_{d}}{2\sqrt{K_{p}M}},\qquad \omega_{n}=\sqrt{\frac{K_{p}}{M}}')
b('$K_{d}$는 점성 마찰 $b$와 같은 역할(가상 댐퍼). $K_{d}>-b,\\ K_{p}>0$이면 안정, 정상상태 오차 0. 임계 감쇠로 고르고, 게인이 너무 크면 채터링·진동, 실제로는 불안정해질 수도 있다(모델링 안 된 동역학, 이산 시간 구현 등).')
sub('중력이 있을 때')
b('PD만 쓰면 정지 상태에서 **정상상태 오차 $=mgr\\cos\\theta/K_{p}$**. 중력을 버틸 토크를 비례항이 내야 하는데, 그러려면 오차가 있어야 하기 때문.')
b('**PID**: 적분항이 정상상태에서 중력을 버티는 토크를 내므로 오차 0. 대신 오차 동역학이 3차가 되고 **$K_{i}$에 상한**이 생긴다(너무 크면 불안정). 과도응답이 나빠질 수 있어(오버슈트·진동) 로봇에서는 $K_{i}$를 0이나 작게 두고 적분값에 상한을 건다.')
sub('모델을 쓰는 제어')
b('**앞먹임만**: $\\tau=\\tilde{M}\\ddot{\\theta}_{d}+\\tilde{h}(\\theta_{d},\\dot{\\theta}_{d})$. 모델이 완벽하지 않으면 오차를 회복하지 못한다.')
b('**Computed torque**: 앞먹임 가속도 + PID 가속도에 모델 $\\tilde{M}$을 곱하고, 중력·마찰 $\\tilde{h}$를 더한다.')
eq(r'\tau=\tilde{M}(\theta)\left(\ddot{\theta}_{d}+K_{p}\theta_{e}+K_{i}\int\theta_{e}\,dt+K_{d}\dot{\theta}_{e}\right)+\tilde{h}(\theta,\dot{\theta})')
b('모델이 정확하면 동역학이 선형화돼 **임의 궤적에서도** 오차 동역학이 안정·선형이 된다. PID만 쓸 때보다 추종이 좋고 제어 노력(토크 제곱 적분)도 적다. 모델이 나쁘면 오히려 손해, 실시간 계산 부담도 있다.')
b('**간단판: PD + 중력 보상** — 전체 동역학 대신 중력 토크 $\\tilde{g}(\\theta)$만 계산해 더한다. 계산이 훨씬 가볍다.')
eq(r'\tau=K_{p}\theta_{e}+K_{d}\dot{\theta}_{e}+\tilde{g}(\theta)')
b('다관절: $\\theta,\\tau$가 벡터, $\\tilde{M}$은 질량행렬, $\\tilde{h}$는 코리올리·중력(·마찰). 작업 공간 버전은 말단 렌치를 구해 관절 토크로 바꾼다.')
eq(r'\mathcal{F}_{b}=\tilde{\Lambda}\left(\dot{\mathcal{V}}_{d}+K_{p}X_{e}+K_{i}\int X_{e}\,dt+K_{d}\mathcal{V}_{e}\right)+\tilde{\eta},\qquad \tau=J_{b}^{T}\mathcal{F}_{b}')
cy('C1 = 위치 제어만(중력 미보상) → 처짐 $=g/K_{p}$만큼 오차. C2 = **PD + 중력 보상**(qfrc_bias) → 오차 0. 다음 단계는 computed torque와 임피던스(현대 파트).')

doc.save(OUT)
print('saved', OUT)
