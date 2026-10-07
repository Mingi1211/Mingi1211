import sys, copy
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

SRC, OUT = 'src.pptx', sys.argv[1]
IMG = {
    'bridge': r'C:\STM32_~2\claude\C--Users-----Desktop------------------\f454b662-091b-4a3c-af4d-fa842e89af86\scratchpad\mj\bridge_fig.png',
    'c1': r'C:\Users\김민기\Desktop\광운대학교\학부연구생\군집 휴머노이드\1주차_실습자료\cyclo_study\C1_bridge.png',
    'c2': r'C:\Users\김민기\Desktop\광운대학교\학부연구생\군집 휴머노이드\1주차_실습자료\cyclo_study\C2_bridge.png',
}
p = Presentation(SRC)
LAYOUT = next(l for l in p.slide_layouts if l.name == '2_Title and Content')
DARK = RGBColor(0x26, 0x26, 0x26)



import re as _re
def math_runs(para, text, size):
    toks = _re.split(r'([_^](?:\{[^}]*\}|.))', text)
    for t in toks:
        if not t:
            continue
        base = 0
        if t[0] in '_^' and len(t) > 1:
            base = -25000 if t[0] == '_' else 30000
            t = t[1:].strip('{}')
        r = para.add_run(); r.text = t
        r.font.size = Pt(size); r.font.name = 'Cambria Math'; r.font.color.rgb = DARK
        if base:
            r.font._rPr.set('baseline', str(base))

def set_title(slide, text):
    slide.shapes.title.text = text


def clear_body(slide):
    """레이아웃의 본문·하단 개체 틀 제거 (제목·슬라이드 번호만 남김)"""
    for sh in list(slide.placeholders):
        if sh.placeholder_format.idx not in (0,) and sh.placeholder_format.type not in (13,):  # 13 = SLIDE_NUMBER
            if 'Slide Number' in sh.name or '슬라이드 번호' in sh.name:
                continue
            sh._element.getparent().remove(sh._element)


def box(slide, x, y, w, h, lines, size=14, header=None):
    """lines: [(text, level, bold)]"""
    if header:
        hb = slide.shapes.add_shape(1, Cm(x), Cm(y), Cm(w), Cm(0.75))
        hb.fill.solid(); hb.fill.fore_color.rgb = RGBColor(0x3B, 0x4A, 0x6B)
        hb.line.fill.background()
        tf = hb.text_frame; tf.text = header
        r = tf.paragraphs[0].runs[0]; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
        y += 0.9; h -= 0.9
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame; tf.word_wrap = True
    for k, item in enumerate(lines):
        para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        para.space_after = Pt(4)
        if item[0] == '=':
            para.alignment = PP_ALIGN.CENTER
            para.space_before = Pt(4); para.space_after = Pt(8)
            math_runs(para, item[1], size)
            continue
        text, lvl, bold = item
        prefix = '• ' if lvl == 0 else '   – '
        r = para.add_run(); r.text = prefix + text if text else ''
        r.font.size = Pt(size if lvl == 0 else size - 2); r.font.bold = bold; r.font.color.rgb = DARK
    return tb


def eqline(slide, x, y, w, text, size=15):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(1.0))
    para = tb.text_frame.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
    math_runs(para, text, size)
    return tb


def new_slide(title):
    s = p.slides.add_slide(LAYOUT)
    set_title(s, title)
    clear_body(s)
    return s


# ---------------------------------------------------------------- 1. 표지 (원본 1번 재사용)
s1 = p.slides[0]
for sh in s1.shapes:
    if not sh.has_text_frame:
        continue
    if sh.name == '제목 4':
        sh.text_frame.paragraphs[0].runs[0].text = '1주차 진행상황 보고'
        for r in sh.text_frame.paragraphs[0].runs[1:]:
            r.text = ''
    elif sh.name == '부제목 5':
        sh.text_frame.paragraphs[0].runs[0].text = 'Cyclo Control × MuJoCo 실습 (AI Worker2)'
        for r in sh.text_frame.paragraphs[0].runs[1:]:
            r.text = ''

# ---------------------------------------------------------------- 2. 이번 주 한 것
s = new_slide('이번 주 한 것')
box(s, 1.2, 2.3, 15.4, 7.6, [
    ('Modern Robotics (Northwestern 강의 + 교재)', 0, True),
    ('3.2.3 지수 좌표 · 5.1 자코비안 · 5.3 특이점', 1, False),
    ('6.2–6.3 수치·역속도 IK · 8.3 RNEA', 1, False),
    ('9.1–9.2 궤적 · 11.3 속도 제어 · 11.4 토크 제어', 1, False),
    ('ROBOTIS Cyclo Control 소스 (MoveL, QP IK)', 0, True),
], header='공부')
box(s, 17.3, 2.3, 15.4, 7.6, [
    ('Ubuntu 24.04 + ROS 2 Jazzy, Cyclo 빌드 (수정 없음)', 0, False),
    ('MuJoCo AI Worker2(FFW-SG2) ↔ Cyclo 브리지 작성', 0, False),
    ('브리지 = 실로봇 follower 역할 (관절 궤적 수신, /joint_states 100 Hz)', 1, False),
    ('실험 C1 기본 이동 · C2 중력 보상 · C4 자기충돌 · C5 도달 불가', 0, True),
], header='셋업 · 개발')
s.shapes.add_picture(IMG['bridge'], Cm(6.5), Cm(9.0), width=Cm(21.0))

# ---------------------------------------------------------------- 3. MoveL 한 주기
s = new_slide('Cyclo MoveL 한 주기 (100 Hz) ↔ MR')
rows = [
    ('① 궤적 생성 — 3차 보간', 'MR 9.2', 'x_{ref}(t),   s(t) = 3(t/T)^2 − 2(t/T)^3,   R(s) = R_0 exp(log(R_0^T R_f) s)'),
    ('② 손끝 목표 속도', 'MR 11.3.3 (11.18)', 'v_d = v_{ff} + K_p e,    e = [ p_{ref} − p ;  log(R_{ref} R^T) ],    K_p = 50'),
    ('③ QP 역속도 IK', 'MR 5.1 · 6.3 + CBF', 'min ‖J q̇ − v_d‖^2_W + q̇^T W_d q̇ + ρ Σs    s.t. 속도 한계, 관절 한계·충돌 CBF'),
    ('④ 적분 → 관절 위치 명령', '—', 'q_d ← q_d + q̇ Δt    (측정값이 아니라 직전 명령값 기준)'),
]
y = 2.6
for name, mr, eq in rows:
    hb = s.shapes.add_shape(1, Cm(1.2), Cm(y), Cm(9.0), Cm(2.6))
    hb.fill.solid(); hb.fill.fore_color.rgb = RGBColor(0x3B, 0x4A, 0x6B); hb.line.fill.background()
    tf = hb.text_frame; tf.word_wrap = True
    tf.text = name
    r = tf.paragraphs[0].runs[0]; r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
    pr = tf.add_paragraph(); rr = pr.add_run(); rr.text = mr; rr.font.size = Pt(11); rr.font.color.rgb = RGBColor(0xDD, 0xE3, 0xF0)
    eqline(s, 10.6, y + 0.8, 22.0, eq, size=14)
    y += 3.3
box(s, 1.2, 15.9, 31.5, 1.6, [('자코비안: Pinocchio LOCAL_WORLD_ALIGNED — 기준점은 손 원점, 축은 world (space·body 둘 다 아님)', 0, False)], size=13)

# ---------------------------------------------------------------- 4. C1 · C2
s = new_slide('C1 기본 이동 · C2 중력 보상')
box(s, 1.2, 2.3, 15.6, 15.2, [
    ('C1: 오른손 → (0.35, −0.20, 0.85), T = 2 s', 0, True),
    ('Cyclo 계산은 목표 도착, MuJoCo 실제 손은 5.6 mm 처짐', 1, False),
    ('처짐은 하위 관절 위치 제어에서 발생', 0, True),
    ('=', 'τ = K_p (q_d − q) − K_d q̇   →   정지 시  q_d − q = g(q) / K_p'),
    ('Cyclo는 직전 명령값으로 적분 → 처짐을 모르고 보정하지 않음', 1, False),
    ('C2: 브리지(하위)에 중력·코리올리 토크 추가', 0, True),
    ('MuJoCo가 RNEA(θ̈ = 0)로 구한 c(q, q̇) + g(q)  (MR 8.3)', 1, False),
    ('=', 'τ = K_p (q_d − q) − K_d q̇ + g̃(q)    (MR 11.4, 식 11.38)'),
    ('5.6 mm → 0.0 mm (Cyclo는 그대로) → 위치 명령 아래 힘 계층 필요', 1, False),
], size=14)
s.shapes.add_picture(IMG['c1'], Cm(17.4), Cm(2.3), height=Cm(7.6))
s.shapes.add_picture(IMG['c2'], Cm(17.4), Cm(10.1), height=Cm(7.6))

# ---------------------------------------------------------------- 5. C4 · C5
s = new_slide('C4 자기충돌 회피 · C5 도달 불가 목표')
box(s, 1.2, 2.3, 15.4, 15.2, [
    ('두 손을 모은 뒤 오른손을 왼손 자리로', 0, False),
    ('오른손: 목표 약 9 cm 앞에서 정지', 1, False),
    ('왼손(명령 없음): 바깥으로 약 8 cm 비켜남', 1, False),
    ('자기충돌 CBF (거리 5 cm 안에서만, 안전거리 2 cm, α = 50)', 0, True),
    ('=', '∇d^T q̇ + s ≥ −α (d − d_{safe})'),
    ('경계에 가까울수록 다가가는 속도 허용치 → 0', 1, False),
    ('양팔을 하나의 QP로 → 기울기에 양팔 관절 포함', 0, True),
    ('"오른손 덜 가기 + 왼손 비키기"가 비용 최소', 1, False),
], header='C4')
box(s, 17.3, 2.3, 15.4, 15.2, [
    ('작업공간 밖 목표 (x 0.80) → x 0.602, z 0.926에서 정지', 0, False),
    ('팔을 끝까지 뻗은 작업공간 경계 = 특이점 근처 (MR 5.3)', 0, True),
    ('목표 추종은 제약이 아니라 비용 → 해가 항상 있음', 0, False),
    ('특이점 전용 제약은 미구현 → 감쇠 항이 담당', 0, True),
    ('관절을 돌려도 손끝이 목표 쪽으로 안 감 → 추종 비용 그대로', 1, False),
    ('감쇠 비용만 증가 → 관절 속도 ≈ 0 → 발산 없이 정지', 1, False),
], header='C5')

# ---------------------------------------------------------------- 6. 정리 · 다음 주
s = new_slide('정리 · 다음 주')
box(s, 1.2, 2.3, 15.4, 15.2, [
    ('Cyclo = 기구학·속도 수준 상위 계층 (질량·중력 모름)', 0, False),
    ('처짐 같은 하위 오차는 하위 계층(중력 보상·임피던스)에서 처리', 0, False),
    ('QP = DLS(감쇠) + 부등식 제약(CBF) + slack', 0, False),
    ('MR 수식은 역할 위주로 정리, 유도는 진행 중', 0, False),
], header='정리')
box(s, 17.3, 2.3, 15.4, 15.2, [
    ('MR 5.1 자코비안 · 6.3 역속도 IK 유도 이어서', 0, False),
    ('C5 자세에서 자코비안 rank · 조작성(5.4) 직접 확인', 0, False),
    ('3차 vs 5차 시간 스케일링 비교', 0, False),
    ('질문', 0, True),
    ('실기에서 팔 모터 전류(토크) 제어 모드 사용 가능 여부', 1, False),
], header='다음 주')

# ---------------------------------------------------------------- 원본 2번 이후 슬라이드 삭제
sldIdLst = p.slides._sldIdLst
ids = list(sldIdLst)
n_new = 5
for sid in ids[1:len(ids) - n_new]:
    rId = sid.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    p.part.drop_rel(rId)
    sldIdLst.remove(sid)
p.save(OUT)
print('saved', OUT, len(p.slides))
