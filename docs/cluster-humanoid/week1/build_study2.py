import sys, os
HERE = os.path.dirname(os.path.abspath(__file__))
_src = open(os.path.join(HERE, '..', 'plan', 'build_manual.py'), encoding='utf-8').read()
exec(_src.split('# =====')[0])          # 바탕체 10pt, 세로 A4, code()/table()/bullet() 등
from docx.shared import Cm as _Cm

SCENE = 'ai_worker/ffw_description/mujoco/ffw_sg2/scene.xml'
SRC = 'source /opt/ros/jazzy/setup.bash && source ~/cyclo_ws/install/setup.bash'


def step(no, title, when=''):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    shade_el(p._p.get_or_add_pPr(), 'D9E2F3')
    run(p, f'☐ STEP {no}. ', bold=True, size=11)
    run(p, title, bold=True, size=11)
    if when:
        run(p, f'   ({when})', size=9)


def ok(text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.4); p.paragraph_format.space_after = Pt(1)
    run(p, '✔ 확인: ', bold=True); rich(p, text)


def ng(text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.4); p.paragraph_format.space_after = Pt(1)
    run(p, '✗ 안 되면: ', bold=True); rich(p, text)


def part(title, sub):
    h = doc.add_heading(level=1); run(h, title, bold=True, size=13)
    para(sub, after=4)


def picture(path, width_cm, caption):
    doc.add_picture(os.path.join(HERE, path), width=_Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(caption, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)


def mr(text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.4); p.paragraph_format.space_after = Pt(1)
    shade_el(p._p.get_or_add_pPr(), 'FFF2CC')
    run(p, '▶ MR 읽기  ', bold=True); rich(p, text)


# =====================================================================
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(t, '1주차 정기미팅 — Cyclo 공부 매뉴얼', bold=True, size=16)
para('AI Worker2(FFW-SG2) 한 대 · 미니 Cyclo로 구조 이해 → 실제 Cyclo를 MuJoCo에서 실행 · Modern Robotics로 개념 정리',
     align=WD_ALIGN_PARAGRAPH.CENTER)
para('군집 휴머노이드 과제 · 김민기 · 보고: 10/8(목) 정기 미팅 · 개정 2026-10-04 (일정 반영, 실제 Cyclo 실행 추가)',
     align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

heading('0. 일정과 결과물')
table([2.6, 2.6, USABLE - 5.2], ['날짜', '시간', '할 일'], [
    ['10/4 (일)', '3시간', '**PART 0 — MR 이론 1**: 4.1 순기구학, 5.1·5.3 자코비안·특이점, 6.2·6.3 수치 IK·역속도 기구학'],
    ['10/5 (월)', '하루 종일', ['**09:00** STEP 1 ROS 2 설치부터 걸어 두기 (다운로드가 길다)',
                              '**09:00~12:30 PART A·B** 환경 → 미니 Cyclo → 실험',
                              '**13:30~17:30 PART C** 실제 Cyclo 빌드 → MuJoCo 브리지 연결 → 실험',
                              '**19:00~22:00 PART D** Cyclo 소스 읽기 + MR 이론 2 (9.1·9.2, 11.3, 8.3 개요)',
                              '**22:00~23:00 PART E** 화요일 공유용 1장 정리']],
    ['10/6 (화)', '오전', '현대와 진행상황 공유 (PART E 1장 사용). 추가 작업 없음'],
    ['10/7 (수)', '3~4시간', '**PART F** 못 한 부분 보완 → 슬라이드 → 예상 질문 리스트로 개념 점검'],
    ['10/8 (목)', '', '정기 미팅 보고'],
], first_bold=True)
para('**미팅까지의 결과물**: ① Cyclo 아키텍처 그림 ② QP 정식화 + MR 대응표 ③ 미니 Cyclo 실험 결과 ④ **실제 Cyclo를 MuJoCo에서 돌린 결과(영상·그래프)** ⑤ 예상 질문과 답')
para('표시: ☐ 끝나면 체크 · ✔ 이렇게 나오면 성공 · ✗ 안 될 때 · ▶ MR 읽기. 작업 폴더는 ~/cyclo_study(MuJoCo·스크립트), ~/cyclo_ws(ROS 2 워크스페이스).', size=9)
para('**같은 폴더의 파일 5개를 Ubuntu ~/cyclo_study 로 옮긴다**: arm_ik_demo.py, model_info.py, cyclo_mujoco_bridge.py, plot_bridge.py, 이 문서.', size=9)

heading('1. 큰 그림 — 두 단계')
bullet('**1단계 (오전)**: Cyclo가 100 Hz마다 하는 계산(목표 자세 → 3차 보간 → 목표 속도 → QP IK → 관절 명령)을 파이썬 150줄 "미니 Cyclo"로 직접 재현해 구조를 이해한다')
bullet('**2단계 (오후)**: 실제 Cyclo(ROBOTIS 코드, 수정 없음)를 빌드하고, 실로봇 대신 MuJoCo를 붙여 같은 명령을 보낸다. 연결은 **cyclo_mujoco_bridge.py** 하나로 한다')
picture('bridge_fig.png', 15.5, '2단계 구성 — 브리지가 실로봇의 follower 역할(관절 궤적 받기, /joint_states 내보내기)을 대신한다')
bullet('브리지 방식을 고른 이유: ros2_control용 MuJoCo 패키지를 쓰면 컨트롤러 이름·토픽 리매핑·부분 관절 명령 문제를 따로 풀어야 한다. 브리지는 Cyclo가 이미 쓰는 토픽 이름을 그대로 받으므로 **Cyclo 쪽은 아무것도 고치지 않는다**')
bullet('작성자가 확인한 범위: 브리지의 MuJoCo 부분(관절 31개 이름이 Cyclo URDF와 모두 일치, 명령 적용, 손끝 기록)은 직접 실행해 확인. **ROS 2 연결 부분은 Windows라 실행해 보지 못했다** → 막히면 부록 A')

# ---------------------------------------------------------------------
part('PART 0 — MR 이론 1 (일, 3시간)', '목표: 내일 코드에서 만날 개념을 미리 잡는다. 각 절마다 노트에 "한 줄 정의 + 식 하나"를 적는다.')
table([2.6, 1.6, 6.4, USABLE - 10.6], ['MR 절', '시간', '핵심', '내일 어디서 만나나'], [
    ['4.1 PoE 순기구학', '40분', 'T(θ) = e^[S1]θ1 ⋯ e^[Sn]θn M. 관절 스크류 축과 홈 자세 M', 'mj_forward, Pinocchio forwardKinematics'],
    ['5.1 자코비안', '70분', 'space 자코비안 J_s(θ)와 body 자코비안 J_b(θ)의 차이, V = J θ̇, 둘 사이 변환 (Ad)', 'mj_jacBody, Pinocchio LOCAL_WORLD_ALIGNED (둘 다 그 중간)'],
    ['5.3 특이점', '20분', 'rank J < 6, 그 근처에서 θ̇가 커지는 이유. (5.4 조작성은 훑기)', 'DLS 감쇠 λ, Cyclo 특이점 slack'],
    ['6.2 수치 IK', '30분', 'Newton–Raphson: θ ← θ + J⁺(θ)·V_err 를 수렴할 때까지 반복', '미니 Cyclo는 주기마다 1번만 → 미분 IK'],
    ['6.3 역속도 기구학', '20분', 'θ̇ = J⁺V, 7자유도(여유 자유도)에서 최소 노름 해, 가중 의사역행렬', 'DLS 식, QP 비용의 댐핑 항'],
], first_bold=True, center_cols=(1,))
para('노트에 꼭 답을 적어 둘 질문 (수요일 예상 질문과 같다):', bold=True)
for q in ['① space 자코비안과 body 자코비안은 각각 어느 좌표계에서 본 twist인가? 7자유도 팔의 자코비안 크기는?',
          '② 특이점 근처에서 의사역행렬 해는 왜 커지나? 감쇠(λ²I)를 더하면 무엇을 얻고 무엇을 잃나?',
          '③ Newton–Raphson IK와 "주기마다 한 번 푸는 미분 IK"의 차이는?']:
    bullet(q)

# ---------------------------------------------------------------------
part('PART A — 환경과 모델 (월 09:00~10:30)', '목표: ROS 2 설치를 걸어 두고, 그동안 MuJoCo에서 AI Worker2를 띄운다.')

step(1, 'ROS 2 Jazzy 설치 걸어 두기 (이미 있으면 건너뜀)', '설치 30~60분, 그동안 STEP 2~4 진행')
para('확인: ls /opt/ros 에 jazzy가 있으면 이미 설치된 것. 없으면 새 터미널에서 (공식 문서: docs.ros.org/en/jazzy/Installation/Ubuntu-Install-Debs.html):')
code(r'''sudo apt install -y software-properties-common curl
sudo add-apt-repository -y universe
export ROS_APT_SOURCE_VERSION=$(curl -s https://api.github.com/repos/ros-infrastructure/ros-apt-source/releases/latest | grep -F "tag_name" | awk -F\" '{print $4}')
curl -L -o /tmp/ros2-apt-source.deb "https://github.com/ros-infrastructure/ros-apt-source/releases/download/${ROS_APT_SOURCE_VERSION}/ros2-apt-source_${ROS_APT_SOURCE_VERSION}.$(. /etc/os-release && echo ${UBUNTU_CODENAME:-${VERSION_CODENAME}})_all.deb"
sudo dpkg -i /tmp/ros2-apt-source.deb
sudo apt update
sudo apt install -y ros-jazzy-desktop ros-dev-tools''')
ok('source /opt/ros/jazzy/setup.bash 후 ros2 --help 가 나온다')
ng('명령이 공식 문서와 다르면 공식 문서를 따른다 (설치 방식이 가끔 바뀐다)')

step(2, 'MuJoCo용 파이썬 가상환경', '10분')
code('''mkdir -p ~/cyclo_study && cd ~/cyclo_study
sudo apt install -y python3-venv git
python3 -m venv venv && source venv/bin/activate
pip install --upgrade pip
pip install mujoco numpy matplotlib osqp scipy''')
ok('python3 -c "import mujoco; print(mujoco.__version__)" → **3.4 이상** (작성자 3.14.0 확인)')
ng('"body mass is too small" → MuJoCo 옛 버전. pip install -U mujoco')

step(3, 'AI Worker 모델 받고 뷰어로 열기', '20분 · 약 150MB')
code('''cd ~/cyclo_study
git clone --depth 1 --filter=blob:none --sparse https://github.com/ROBOTIS-GIT/ai_worker.git
cd ai_worker && git sparse-checkout set ffw_description && cd ..
python3 -m mujoco.viewer --mjcf=''' + SCENE)
ok('로봇이 바닥에 서 있다. 오른쪽 Control 패널에서 arm_r_joint1~7 슬라이더를 움직여 각 관절의 회전 방향을 본다')
ng('SSH라 화면이 없으면 뷰어는 건너뛰고, 이후 스크립트는 --view / 뷰어 없이 실행')

step(4, '모델 정보 확인', '10분')
code('python3 model_info.py')
ok('nq = 38, nv = 37, nu = 25, timestep = 0.002 · arm_r_joint1 kp = 3000, 힘 범위 ±61.4 · 손끝 [−0.02, −0.227, 0.725]')
bullet('nq ≠ nv: 베이스가 free joint(위치는 쿼터니언 4개, 속도는 각속도 3개). 관절마다 **위치 액추에이터**(kp·오차 − kv·속도) = DYNAMIXEL 위치 제어 흉내')

# ---------------------------------------------------------------------
part('PART B — 미니 Cyclo로 구조 이해 (월 10:30~12:30)', '목표: Cyclo MoveL의 계산을 직접 돌려 보고, 무엇이 왜 필요한지 숫자로 본다.')

step(5, '기본 실행', '10분')
code('python3 arm_ik_demo.py --scene ' + SCENE + ' --view')
ok('"최대 위치 오차 6.8 mm · 최종 0.1 mm · 최대 관절 속도 1.06 rad/s" 근처 (작성자 Windows 값), arm_ik_plot.png 생성')
picture('E1_plot.png', 11.5, '참고 — 기본 실행 (DLS, 5차, T = 3 s). 목표는 base_link 기준 (0.35, −0.20, 0.85) = Cyclo MoveL 예시와 같다')

step(6, '코드 블록을 Cyclo·MR과 연결', '30분')
table([4.2, 5.6, USABLE - 9.8], ['arm_ik_demo.py', '실제 Cyclo', 'MR'], [
    ['scaling(): 5차 / 3차', 'movel 노드 cubicVector, rotationCubic (Cyclo = **3차**)', '9.2 시간 스케일링'],
    ['p_des = p0 + s·(p1 − p0)', '직선 보간 (MoveL)', '9.2.1 직선 경로'],
    ['xd = v_ff + K·e', 'desired_vel = feedforward + kp_position·error', '11.3 속도 입력 작업공간 제어'],
    ['mj_jacBody → J (6×7)', 'Pinocchio getJacobian, LOCAL_WORLD_ALIGNED', '5.1 자코비안'],
    ['solve_dls()', '(Cyclo에는 없음)', '6.3 역속도, 특이점 감쇠'],
    ['solve_qp()', 'VRController setCost / Bound / Ineq', 'MR 밖 — QP, CBF (STEP 15)'],
    ['q_cmd = q_base + dq·Δt', 'q_desired_ = q_feedback + optimal_velocities·time_step', '수치 적분'],
], first_bold=False)

step(7, '실험 — 꼭 할 것 5개 (★), 시간 나면 나머지', '60분')
code('python3 arm_ik_demo.py --scene ' + SCENE + ' --T 1.0 --solver qp --out E7      # 예시')
table([1.0, 4.6, 1.9, 1.9, 1.9, USABLE - 11.3], ['#', '옵션', '최대 오차', '최종 오차', '최대 q̇', '볼 것 · 내 결과'], [
    ['E1★', '(기본) DLS, 5차, T=3, K=50', '6.8 mm', '0.1 mm', '1.06', '기준'],
    ['E2', '--profile cubic', '7.2', '0.0', '0.98', 'Cyclo와 같은 3차'],
    ['E3', '--K 5', '13.2', '3.5', '1.07', '피드백 게인 작으면 뒤처짐'],
    ['E4', '--lam 0.3', '38.5', '0.0', '1.41', '감쇠 크면 느림'],
    ['E5★', '--solver qp', '4.7', '0.0', '1.11', 'Cyclo식 QP'],
    ['E6★', '--T 1.0', '180.4', '112.3', '12.92', 'DLS: 속도 한계 없음 → 폭주'],
    ['E7★', '--T 1.0 --solver qp', '104.2', '3.6', '1.50', 'QP: 한계 지킴, 대신 뒤처짐'],
    ['E8★', '--integrate feedback', '32.1', '11.8', '3.58', 'Cyclo 적분 방식 → 중력 처짐'],
    ['E9', '--integrate feedback --gravcomp', '28.6', '3.4', '3.24', '중력 보상 feedforward → 처짐 감소'],
], first_bold=True, center_cols=(0, 2, 3, 4))
picture('compare.png', 14.5, '왼쪽: 명령 기준 적분(E1) vs Cyclo식 측정값 기준 적분(E8) / 오른쪽: 빠른 이동에서 DLS(E6) vs QP(E7) 관절 속도')
bullet('**E6 vs E7**: 제약 없는 DLS는 관절 속도가 12.9 rad/s까지 튄다. QP는 1.5 rad/s를 넘지 않는다 → **Cyclo가 QP를 쓰는 이유**')
bullet('**E8 → E9**: Cyclo처럼 "측정 관절각 + q̇Δt"를 명령으로 주면 위치 액추에이터 목표가 측정값과 거의 같아 중력을 버틸 힘이 모자라 11.8 mm가 남는다. '
       '팔에 중력·코리올리 토크를 feedforward로 더하면(--gravcomp, RNEA로 구하는 값) 3.4 mm로 준다 → **교수님이 말씀하신 "처짐을 모델에 넣어 feedforward"**')

# ---------------------------------------------------------------------
part('PART C — 실제 Cyclo를 MuJoCo에서 실행 (월 13:30~17:30)', '목표: ROBOTIS의 Cyclo를 수정 없이 빌드하고, 브리지로 MuJoCo 로봇을 움직인다. 미팅의 핵심 시연.')

step(8, 'Cyclo 빌드', '30~40분')
code('''source /opt/ros/jazzy/setup.bash
mkdir -p ~/cyclo_ws/src && cd ~/cyclo_ws/src
git clone https://github.com/ROBOTIS-GIT/cyclo_control.git
vcs import . < cyclo_control/cyclo_control_ci.repos     # robotis_interfaces를 같이 받음
cd ~/cyclo_ws
sudo rosdep init            # 처음 한 번만 (이미 했다는 에러는 무시)
rosdep update
rosdep install --from-paths src --ignore-src -r -y --rosdistro jazzy
colcon build --symlink-install --cmake-args -DCMAKE_BUILD_TYPE=Release
source install/setup.bash''')
ok('colcon이 Summary: N packages finished 로 끝나고 ros2 interface show robotis_interfaces/msg/MoveL 이 나온다')
ng('Pinocchio를 못 찾으면 sudo apt install ros-jazzy-pinocchio 후 다시 빌드. numpy 2 관련 에러면 pip로 깐 numpy가 섞인 것 → 이 터미널에서는 venv를 끈다(deactivate)')

step(9, '브리지용 파이썬 환경', '10분')
para('브리지는 rclpy(ROS)와 mujoco(pip)를 같이 써야 한다. ROS 파이썬을 보이게 한 venv를 하나 더 만든다:')
code('''cd ~/cyclo_study
python3 -m venv --system-site-packages venv_ros
source venv_ros/bin/activate
pip install mujoco "numpy<2" matplotlib''')
ok(SRC + ' 후 venv_ros를 켜고 python3 -c "import rclpy, mujoco" 가 에러 없이 끝난다')
bullet('모든 터미널에서 같은 컴퓨터 안에서만 통신하게 하면 랩 네트워크의 다른 ROS와 섞이지 않는다: export ROS_AUTOMATIC_DISCOVERY_RANGE=LOCALHOST')

step(10, '세 터미널로 실행', '20분')
para('**세 터미널 모두 먼저** 아래를 입력한다:')
code(SRC + '\nexport ROS_AUTOMATIC_DISCOVERY_RANGE=LOCALHOST')
table([1.6, 4.0, USABLE - 5.6], ['터미널', '역할', '명령'], [
    ['T1', 'MuJoCo 브리지', ['source ~/cyclo_study/venv_ros/bin/activate && cd ~/cyclo_study', 'python3 cyclo_mujoco_bridge.py --scene ' + SCENE]],
    ['T2', '실제 Cyclo', ['ros2 launch cyclo_motion_controller_ros ai_worker_controller.launch.py controller_type:=movel']],
    ['T3', '명령·확인', ['ros2 topic hz /joint_states          # 약 100 Hz', 'ros2 topic echo --once /r_gripper_pose   # Cyclo가 계산한 손끝']],
], first_bold=True, center_cols=(0,))
para('T3에서 Cyclo README와 같은 명령을 보낸다:')
code('''ros2 topic pub --once /r_goal_move robotis_interfaces/msg/MoveL "{pose: {header: {frame_id: 'base_link'}, pose: {position: {x: 0.35, y: -0.20, z: 0.85}, orientation: {x: 0.0, y: 0.0, z: 0.0, w: 1.0}}}, time_from_start: {sec: 2, nanosec: 0}}"''')
ok('MuJoCo 창에서 오른팔이 2초 동안 앞으로 올라간다. **화면 녹화 (미팅 시연 영상)**')
ok('T1을 Ctrl+C로 끄면 bridge_log.csv가 남는다 → STEP 12에서 그래프')
ng('팔이 안 움직이면: T3에서 ros2 topic echo --once /leader/joint_trajectory_command_broadcaster_right/joint_trajectory → 나오면 브리지 쪽, 안 나오면 Cyclo 쪽(T2 로그에서 joint state 관련 경고 확인). 부록 A')

step(11, '실제 Cyclo 실험', '90분')
para('실험마다 T1 브리지를 새로 켜고(로그 파일 이름 --log C1.csv 처럼 바꿔서), 명령 하나를 보낸 뒤 5초 기다려 끈다. **예측**은 미니 Cyclo 결과로 미리 적어 둔 것이다.')
table([1.0, 6.4, 4.0, USABLE - 11.4], ['#', '방법', '예측', '내 결과'], [
    ['C1★', '기본: 위 명령 (T = 2 s)', '목표 근처에 서지만 중력 처짐으로 수 mm~10 mm 남음 (E8)', ''],
    ['C2★', 'T1을 --gravcomp 로 켜고 C1 반복', '남는 오차 감소 (E9: 11.8 → 3.4 mm)', ''],
    ['C3', 'time_from_start: {sec: 0} (보간 없이 바로 목표)', '빠르게 따라가지만 속도 한계(QP)에 걸려 직선이 아님', ''],
    ['C4★', '자기충돌: ① /l_goal_move (0.35, 0.10, 0.85), /r_goal_move (0.35, −0.10, 0.85)로 두 손을 앞에 모은다 ② /r_goal_move (0.35, 0.10, 0.85) — 오른손을 왼손 자리로', '오른손이 왼손과 일정 거리(약 2 cm + 그리퍼 크기)에서 멈추고 목표까지 못 감 (CBF)', ''],
    ['C5', '도달 불가: position {x: 0.80, y: -0.20, z: 0.85}', '팔을 최대한 뻗고 멈춤. 발산하지 않음 (slack)', ''],
    ['C6', '파라미터: 설정 파일을 복사해 kp_position 50 → 10 또는 cbf_alpha 50 → 5 로 바꾸고 config_file:=<경로> 로 실행', 'kp↓ → 느려짐 / α↓ → 경계에서 더 일찍 감속', ''],
    ['C7', 'FK 일치: T3에서 /r_gripper_pose(Cyclo의 Pinocchio FK) vs 브리지 CSV 마지막 줄(MuJoCo)', 'URDF와 MJCF가 같은 로봇이면 거의 같음', ''],
], first_bold=True, center_cols=(0,))
code('''# C6용 설정 파일 복사
cp ~/cyclo_ws/install/cyclo_motion_controller_ros/share/cyclo_motion_controller_ros/config/ai_worker_config.yaml ~/cyclo_study/my_config.yaml
ros2 launch cyclo_motion_controller_ros ai_worker_controller.launch.py controller_type:=movel config_file:=$HOME/cyclo_study/my_config.yaml''')

step(12, '결과 그래프', '20분')
code('''python3 plot_bridge.py --log C1.csv --target 0.35 -0.20 0.85 --start 3 --out C1
python3 plot_bridge.py --log C2.csv --target 0.35 -0.20 0.85 --start 3 --out C2''')
ok('C1_bridge.png, C2_bridge.png — 위: 손끝 x·y·z와 목표(점선), 아래: 목표까지 거리. --start는 명령을 보낸 시각 근처로 맞춘다')

# ---------------------------------------------------------------------
part('PART D — Cyclo 소스 읽기 + MR 이론 2 (월 19:00~22:00)', '목표: 오늘 돌린 것을 그림과 수식으로 설명할 수 있게 만든다.')
step(13, 'MR 이론 2', '70분')
table([2.8, 1.6, 6.2, USABLE - 10.6], ['MR 절', '시간', '핵심', '연결'], [
    ['9.1~9.2 궤적 생성', '35분', '경로(path)와 시간 스케일링 s(t) 분리. 3차: 시작·끝 속도 0 / 5차: 가속도까지 0', 'Cyclo 3차 보간, E2'],
    ['11.3 속도 입력 제어', '20분', 'V = [Ad]V_d + K_p·X_err : 피드포워드 + 오차 피드백', 'v_d = v_ff + kp·e'],
    ['8.3 RNEA (개요)', '15분', '역동역학 τ = M(θ)θ̈ + c(θ,θ̇) + g(θ) 를 재귀로 계산', '--gravcomp, 다음 단계 힘 계층'],
], first_bold=True, center_cols=(1,))

step(14, 'MoveL 아키텍처 따라가기', '40분')
picture('cyclo_arch.png', 15.5, 'Cyclo MoveL 한 주기 (100 Hz) — 작성자가 소스에서 확인한 구조')
for txt in ['① ai_worker_config.yaml: right_movel_topic = /r_goal_move',
            '② movel 노드: cubicVector / rotationCubic 호출부, desired_vel = feedforward + kp_position·error',
            '③ kinematics_solver.cpp: computeFrameJacobian(..., LOCAL_WORLD_ALIGNED), computeDistances(자기충돌)',
            '④ vr_controller.cpp: setCost(), setBoundConstraint(), setIneqConstraint()',
            '⑤ movel 노드: q_desired_ = q_feedback + optimal_velocities * time_step_',
            '⑥ publishTrajectory(): 오른팔 7 + gripper_r_joint1, time_from_start = trajectory_time(0.0)',
            '⑦ joint_state_timeout: /joint_states가 끊기면 정지·유지']:
    bullet('☐ ' + txt)

step(15, 'QP 정식화', '50분')
table([3.0, 6.6, USABLE - 9.6], ['구성', '수식', '코드 · 의미'], [
    ['결정 변수', 'x = [q̇, s_qmin, s_qmax, s_sing, s_col]', 'q̇ 관절 속도, s slack'],
    ['비용', 'Σ (J_i q̇ − v_d,i)ᵀ W_i (J_i q̇ − v_d,i) + q̇ᵀ W_d q̇ + ρ Σ s', 'setCost(). W = 위치 10 / 자세 1, W_d = 0.1, ρ = 1000'],
    ['속도 한계', 'q̇_min ≤ q̇ ≤ q̇_max,  s ≥ 0', 'setBoundConstraint()'],
    ['관절 한계 CBF', 'q̇ + s ≥ −α (q − q_min),  −q̇ + s ≥ −α (q_max − q)', 'setIneqConstraint(), α = cbf_alpha 50'],
    ['자기충돌 CBF', 'ḋ(q̇) + s ≥ −α (d − d_safe)  (d ≤ buffer일 때)', 'd_safe 0.02, buffer 0.05 → C4 실험'],
    ['특이점', '제약 1개 + slack 1개', '코드로 직접 확인 (조작성 관련으로 추정)'],
], first_bold=True)
bullet('**CBF**: 안전 거리 h ≥ 0을 지키려면 ḣ ≥ −α·h. 경계에 가까울수록 다가가는 속도도 0에 가까워져 넘지 않는다 (MR 밖 개념)')
bullet('**slack**: 제약끼리 충돌해 QP가 안 풀리면 로봇이 멈춘다. 큰 벌점으로 평소엔 0, 불가능할 때만 조금 어긴다 → C5')
bullet('**DLS와의 관계**: 제약을 빼면 해는 (JᵀWJ + W_d)q̇ = JᵀW v_d. W = I, W_d = λ²I이면 DLS. **DLS는 제약 없는 QP의 특수한 경우**')
bullet('**여유 자유도**: 7자유도 팔에서 남는 1자유도는 댐핑 항(속도 크기 최소)이 정한다. 자세 유지 같은 별도 null-space 작업은 없다')

# ---------------------------------------------------------------------
part('PART E — 화요일 공유용 1장 (월 22:00~23:00)', '현대와 오전에 나눌 내용. 그대로 채워 한 장으로 출력한다.')
table([3.6, USABLE - 3.6], ['항목', '채울 것'], [
    ['한 일', '미니 Cyclo로 구조 재현 / 실제 Cyclo 빌드 + MuJoCo 연결 / 실험 C1~C7 중 한 것'],
    ['수치', 'C1 vs C2 남은 오차, E6 vs E7 관절 속도, C4 멈춘 거리'],
    ['막힌 것', '설치·연결에서 막힌 곳과 해결 여부'],
    ['참빛과 연결', '**bimanual_movel = 가상 물체를 잡은 채 양팔 상대 자세 유지** → 참빛 협동 운반과 같은 문제. 참빛 IK를 직접 만들지, Cyclo를 기준선으로 비교할지 같이 정하기'],
    ['현대 몫과 연결', 'Cyclo 출력은 위치 수준 → 현대의 임피던스(전류 제어)는 그 아래 계층. E9의 중력 feedforward가 양쪽 공통 출발점'],
], first_bold=True)

# ---------------------------------------------------------------------
part('PART F — 수요일 최종 준비 (3~4시간)', '못 한 실험 보완 → 슬라이드 → 예상 질문으로 개념 점검.')
heading('슬라이드 8장', 2)
table([1.0, 4.0, USABLE - 5.0], ['#', '제목', '넣을 것'], [
    ['1', '목표', '교수님 조언(IK 구조 이해 우선) → 방법 2단계 (미니 Cyclo → 실제 Cyclo + MuJoCo)'],
    ['2', 'Cyclo 아키텍처', 'STEP 14 그림'],
    ['3', 'IK = QP', 'STEP 15 수식, CBF·slack 한 줄씩, "DLS는 제약 없는 QP"'],
    ['4', 'MR 개념 대응', 'STEP 6 표 + 자코비안 좌표계(space / body / LOCAL_WORLD_ALIGNED)'],
    ['5', '미니 Cyclo 실험', 'E6 vs E7 (QP를 쓰는 이유), E8 → E9 (중력 처짐과 feedforward)'],
    ['6', '실제 Cyclo + MuJoCo', 'STEP 1의 구성 그림, **시연 영상**, C1 vs C2 그래프, C4 충돌 회피'],
    ['7', '관찰과 한계', '시뮬 위치 액추에이터 ≠ 실제 DYNAMIXEL, 고정 궤적(3차)의 한계, 힘 계층 부재'],
    ['8', '다음 주', '실로봇에서 같은 명령 비교, bimanual_movel, 힘 계층(중력 보상 → 임피던스) 위치, 질문 1~2개'],
], first_bold=True, center_cols=(0,))
heading('예상 질문 리스트 (초안 — 수요일에 내 말로 답을 다시 쓴다)', 2)
table([0.8, 6.2, USABLE - 7.0], ['#', '질문', '답의 뼈대'], [
    ['1', 'space / body 자코비안 차이? Cyclo는 어느 것?', 'space = 고정 좌표계 원점에서 본 twist, body = 말단 좌표계. Cyclo는 LOCAL_WORLD_ALIGNED: 말단 원점의 속도를 world 축으로 → 둘 다 아님'],
    ['2', '왜 Newton–Raphson이 아니라 미분 IK?', '100 Hz마다 현재 자세에서 한 번 선형화해 속도를 푼다. 반복 수렴이 필요 없고 제약을 QP로 넣기 쉽다'],
    ['3', '특이점에서 무슨 일? 어떻게 막나?', 'J의 rank가 떨어져 q̇가 커짐. DLS 감쇠 / QP 댐핑 항 / Cyclo 특이점 slack'],
    ['4', '7자유도의 남는 1자유도는?', 'QP 댐핑 항(최소 속도)이 정함. 자세 유지 null-space 작업은 없음 → 확장 포인트'],
    ['5', '3차 vs 5차?', '경계 조건 4개 vs 6개. 5차는 가속도도 0이라 부드러움. Cyclo는 3차, 실험 차이는 작았음'],
    ['6', 'v_d = v_ff + kp·e 의미?', 'MR 11.3. 피드포워드로 따라가고 오차는 피드백으로 줄임. kp가 작으면 뒤처짐(E3)'],
    ['7', 'QP 비용·제약을 말해보라', 'STEP 15 표. OSQP 표준형 ½xᵀPx + qᵀx, l ≤ Ax ≤ u'],
    ['8', 'CBF가 뭔가? α는?', 'ḣ ≥ −α h. α가 크면 경계 가까이까지 빠르게, 작으면 일찍 감속 (C6)'],
    ['9', 'slack은 왜?', '제약이 서로 충돌해도 QP가 풀리게. 벌점 1000 → 평소 0 (C5)'],
    ['10', '왜 측정값 기준으로 적분하나?', '명령 기준이면 로봇이 못 따라갈 때 명령이 폭주(E6, windup). 대신 위치 제어가 약하면 중력 처짐(E8)'],
    ['11', '힘 대응은 어디에 붙이나?', 'Cyclo 출력은 위치. 아래에 중력 feedforward(E9, RNEA) → 전류 제어 모드 임피던스'],
    ['12', '시뮬과 실기 차이는?', 'MuJoCo 위치 액추에이터 kp·kv vs DYNAMIXEL 내부 PID, 통신 지연, 모델 오차(C7)'],
    ['13', '이기종(휴머노이드·모바일 매니퓰레이터)으로 가면?', 'URDF·관절 그룹 교체, 베이스·허리 task 추가, floating base면 접촉 제약 필요 → Cyclo 구조(작업 목록 + 제약) 그대로 확장'],
], first_bold=False, center_cols=(0,))

# ---------------------------------------------------------------------
heading('부록 A. 문제가 생겼을 때')
table([5.6, USABLE - 5.6], ['증상', '해결'], [
    ['body mass is too small', 'MuJoCo 업데이트 (pip install -U mujoco)'],
    ['externally-managed-environment', 'venv를 안 켬'],
    ['mesh file not found', 'STEP 3 sparse checkout 확인 (ffw_description/meshes)'],
    ['rosdep: cannot find pinocchio', 'sudo apt install ros-jazzy-pinocchio'],
    ['브리지: No module named rclpy', 'T1에서 source /opt/ros/jazzy/setup.bash를 venv 켜기 전에 했는지, venv_ros를 --system-site-packages로 만들었는지'],
    ['브리지: numpy 관련 에러', 'venv_ros 안에서 pip install "numpy<2"'],
    ['Cyclo가 명령을 안 냄', 'ros2 topic hz /joint_states 확인. T2 로그 확인. 모든 터미널의 ROS_DOMAIN_ID·ROS_AUTOMATIC_DISCOVERY_RANGE가 같은지'],
    ['궤적은 나오는데 팔이 안 움직임', '브리지가 그 토픽을 받는지 (ros2 topic info -v …/joint_trajectory 에서 구독자 확인)'],
    ['로봇이 넘어지거나 흔들림', '브리지를 다시 켠다. 베이스는 바퀴로 서 있는 free body라 큰 동작에서 흔들릴 수 있다'],
    ['그래도 안 됨', '실제 Cyclo 단계는 다음 주로 넘기고, 미니 Cyclo 결과 + 빌드까지의 로그로 보고 (막힌 지점을 정확히 보고하는 것도 성과)'],
], first_bold=False)

heading('부록 B. 코드 (파일로 같은 폴더에 있음, 읽기용)')
for fn in ['cyclo_mujoco_bridge.py', 'plot_bridge.py', 'arm_ik_demo.py']:
    para(fn, bold=True)
    code(open(os.path.join(HERE, fn), encoding='utf-8').read())

heading('참고', 2)
for s in ['Modern Robotics (Lynch & Park): 4.1, 5.1, 5.3, 6.2~6.3, 8.3, 9.1~9.2, 11.3',
          'Cyclo Control: https://github.com/ROBOTIS-GIT/cyclo_control (README의 설치·실행 방법)',
          'AI Worker MuJoCo 모델: https://github.com/ROBOTIS-GIT/ai_worker (ffw_description/mujoco)',
          'ROS 2 Jazzy 설치: https://docs.ros.org/en/jazzy/Installation/Ubuntu-Install-Debs.html']:
    bullet(s)

doc.save(OUT)
print('saved', OUT)
